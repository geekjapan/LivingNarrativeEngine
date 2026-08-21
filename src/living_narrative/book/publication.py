"""Reproducible reader publication adapters backed by the accepted Markdown manuscript."""

from __future__ import annotations

import hashlib
import html
import os
from pathlib import Path
from typing import Literal

import yaml
from docx import Document
from ebooklib import epub
from pydantic import BaseModel
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen.canvas import Canvas

from living_narrative.book.exporter import export_accepted_manuscript
from living_narrative.state.diff import fsync_directory
from living_narrative.workspace.loader import WorkspacePaths

PublicationFormat = Literal["docx", "epub", "pdf"]


class PublicationVerificationError(ValueError):
    """A publication manifest or its referenced reader artifact is inconsistent."""


class PublicationVerificationResult(BaseModel):
    """Reader-safe verification evidence for one publication manifest."""

    manifest_path: Path
    manuscript_sha256: str
    format_names: list[str]


class PublicationExportResult(BaseModel):
    """Publication files generated from one accepted manuscript generation."""

    format_paths: dict[PublicationFormat, Path]
    manifest_path: Path
    manuscript_sha256: str


def _atomic_write_bytes(path: Path, payload: bytes) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        with temporary.open("wb") as stream:
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        fsync_directory(path.parent)
    finally:
        temporary.unlink(missing_ok=True)


def _markdown_blocks(manuscript: str) -> list[tuple[str, str]]:
    """Reduce canonical reader Markdown to a small deterministic publication block model."""
    blocks: list[tuple[str, str]] = []
    for raw in manuscript.split("\n\n"):
        value = raw.strip()
        if not value:
            continue
        if value.startswith("# "):
            blocks.append(("heading", value[2:].strip()))
        else:
            blocks.append(("paragraph", value))
    return blocks


def _write_docx(manuscript: str, path: Path) -> None:
    document = Document()
    for kind, value in _markdown_blocks(manuscript):
        if kind == "heading":
            document.add_heading(value, level=1)
        else:
            document.add_paragraph(value)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        document.save(temporary)
        _atomic_write_bytes(path, temporary.read_bytes())
    finally:
        temporary.unlink(missing_ok=True)


def _write_epub(manuscript: str, manuscript_sha256: str, path: Path) -> None:
    book = epub.EpubBook()
    book.set_identifier(manuscript_sha256)
    book.set_title("Living Narrative Manuscript")
    book.set_language("ja")
    body = "".join(
        f"<h1>{html.escape(value)}</h1>" if kind == "heading" else f"<p>{html.escape(value)}</p>"
        for kind, value in _markdown_blocks(manuscript)
    )
    chapter = epub.EpubHtml(title="Manuscript", file_name="manuscript.xhtml", lang="ja")
    chapter.content = f"<html><body>{body}</body></html>"
    book.add_item(chapter)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        epub.write_epub(str(temporary), book, {"epub3_landmark": False})
        _atomic_write_bytes(path, temporary.read_bytes())
    finally:
        temporary.unlink(missing_ok=True)


def _write_pdf(manuscript: str, path: Path) -> None:
    """Typeset Japanese reader text with a built-in CID font and deterministic page geometry."""
    font_name = "HeiseiMin-W3"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        canvas = Canvas(str(temporary), pagesize=A4)
        width, height = A4
        y = height - 62
        for kind, value in _markdown_blocks(manuscript):
            size = 16 if kind == "heading" else 11
            leading = 24 if kind == "heading" else 18
            canvas.setFont(font_name, size)
            lines = [value[index : index + 34] for index in range(0, len(value), 34)] or [""]
            for line in lines:
                if y < 62:
                    canvas.showPage()
                    canvas.setFont(font_name, size)
                    y = height - 62
                canvas.drawString(62, y, line)
                y -= leading
            y -= 8
        canvas.save()
        _atomic_write_bytes(path, temporary.read_bytes())
    finally:
        temporary.unlink(missing_ok=True)


def _sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def verify_publication(manifest_path: Path) -> PublicationVerificationResult:
    """Fail closed unless all manifest-linked publication files exist and match their hashes."""
    try:
        manifest = yaml.safe_load(manifest_path.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as exc:
        raise PublicationVerificationError("publication manifest cannot be read") from exc
    if not isinstance(manifest, dict) or manifest.get("schema_version") != 1:
        raise PublicationVerificationError("unsupported publication manifest")
    output_dir = manifest_path.parent
    manuscript = output_dir / "manuscript.md"
    expected_manuscript = manifest.get("manuscript_sha256")
    if not isinstance(expected_manuscript, str) or not manuscript.is_file():
        raise PublicationVerificationError("publication manuscript is missing")
    if _sha256_file(manuscript) != expected_manuscript:
        raise PublicationVerificationError("manuscript sha256 mismatch")
    source_manifest = output_dir / "manuscript_manifest.yaml"
    expected_source = manifest.get("source_manuscript_manifest_sha256")
    if not isinstance(expected_source, str) or not source_manifest.is_file():
        raise PublicationVerificationError("source manuscript manifest is missing")
    if _sha256_file(source_manifest) != expected_source:
        raise PublicationVerificationError("source manuscript manifest sha256 mismatch")
    formats = manifest.get("formats")
    if not isinstance(formats, dict) or set(formats) != {"docx", "epub", "pdf"}:
        raise PublicationVerificationError("publication formats are incomplete")
    for name, artifact in formats.items():
        if not isinstance(artifact, dict):
            raise PublicationVerificationError(f"{name} artifact is invalid")
        filename = artifact.get("filename")
        expected = artifact.get("sha256")
        valid_filename = isinstance(filename, str) and Path(filename).name == filename
        if not valid_filename or not isinstance(expected, str):
            raise PublicationVerificationError(f"{name} artifact is invalid")
        path = output_dir / filename
        if not path.is_file():
            raise PublicationVerificationError(f"{name} artifact is missing")
        if _sha256_file(path) != expected:
            raise PublicationVerificationError(f"{name} sha256 mismatch")
    return PublicationVerificationResult(
        manifest_path=manifest_path,
        manuscript_sha256=expected_manuscript,
        format_names=sorted(formats),
    )


def export_publication(
    workspace: Path | WorkspacePaths, output_dir: Path
) -> PublicationExportResult:
    """Generate DOCX, EPUB, PDF, and one hash-linked manifest from accepted Markdown only.

    The existing Markdown exporter validates every accepted pointer and candidate hash before this
    adapter work begins. The publication manifest is published last, so consumers never see a
    complete publication claim when any format adapter fails.
    """
    manuscript_export = export_accepted_manuscript(workspace, output_dir)
    manuscript = manuscript_export.manuscript_path.read_text(encoding="utf-8")
    manuscript_bytes = manuscript.encode("utf-8")
    manuscript_sha256 = hashlib.sha256(manuscript_bytes).hexdigest()
    generated: dict[PublicationFormat, Path] = {
        "docx": output_dir / "manuscript.docx",
        "epub": output_dir / "manuscript.epub",
        "pdf": output_dir / "manuscript.pdf",
    }
    _write_docx(manuscript, generated["docx"])
    _write_epub(manuscript, manuscript_sha256, generated["epub"])
    _write_pdf(manuscript, generated["pdf"])

    source_manifest = manuscript_export.manifest_path.read_bytes()
    publication_manifest = {
        "schema_version": 1,
        "manuscript_sha256": manuscript_sha256,
        "source_manuscript_manifest_sha256": hashlib.sha256(source_manifest).hexdigest(),
        "quality_gate": "accepted_chapters_only",
        "license": "unspecified",
        "formats": {
            name: {
                "filename": path.name,
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
            for name, path in generated.items()
        },
    }
    manifest_path = output_dir / "publication_manifest.yaml"
    _atomic_write_bytes(
        manifest_path,
        yaml.safe_dump(publication_manifest, allow_unicode=True, sort_keys=False).encode("utf-8"),
    )
    return PublicationExportResult(
        format_paths=generated,
        manifest_path=manifest_path,
        manuscript_sha256=manuscript_sha256,
    )
